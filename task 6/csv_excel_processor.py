# csv_excel_processor.py

import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule


# ==========================================
# CSV PROCESSOR
# ==========================================

class CSVProcessor:

    @staticmethod
    def read_csv(file_path):

        try:
            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:

                return list(
                    csv.DictReader(file)
                )

        except FileNotFoundError:

            print("CSV file not found.")
            return []

    @staticmethod
    def write_csv(file_path, data):

        if not data:
            print("No data available.")
            return

        with open(
            file_path,
            "w",
            newline="",
            encoding="utf-8"
        ) as file:

            writer = csv.DictWriter(
                file,
                fieldnames=data[0].keys()
            )

            writer.writeheader()
            writer.writerows(data)

        print("CSV written successfully.")

    @staticmethod
    def merge_csv_files(
        input_files,
        output_file
    ):

        merged_data = []

        for file_name in input_files:

            merged_data.extend(
                CSVProcessor.read_csv(file_name)
            )

        CSVProcessor.write_csv(
            output_file,
            merged_data
        )

        print("CSV files merged successfully.")

    @staticmethod
    def remove_duplicates(data, key):

        unique = {}

        for row in data:
            unique[row[key]] = row

        return list(unique.values())

    @staticmethod
    def validate_data(
        data,
        required_columns
    ):

        errors = []

        for row_no, row in enumerate(
            data,
            start=1
        ):

            for column in required_columns:

                if (
                    column not in row
                    or row[column] == ""
                ):

                    errors.append(
                        f"Row {row_no}: "
                        f"{column} missing"
                    )

        return errors


# ==========================================
# EXCEL PROCESSOR
# ==========================================

class ExcelProcessor:

    @staticmethod
    def create_workbook(file_name):

        wb = Workbook()

        # Sheet 1
        ws = wb.active
        ws.title = "Sales"

        headers = [
            "Product",
            "Quantity",
            "Price",
            "Revenue"
        ]

        ws.append(headers)

        data = [

            ["Laptop", 10, 50000],
            ["Phone", 15, 20000],
            ["Tablet", 8, 15000],
            ["Monitor", 12, 10000]

        ]

        for row in data:
            ws.append(row)

        # Formula Generation
        for row in range(2, 6):

            ws[f"D{row}"] = (
                f"=B{row}*C{row}"
            )

        # Header Formatting
        for cell in wscell.font : Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                start_color="4F81BD"
            )

        # Multiple Sheets
        summary = wb.create_sheet(
            title="Summary"
        )

        summary["A1"] = (
            "Total Revenue"
        )

        summary["B1"] = (
            "=SUM(Sales!D2:D5)"
        )

        # Bar Chart
        chart = BarChart()
        chart.title = "Revenue Analysis"

        data_ref = Reference(
            ws,
            min_col=4,
            min_row=1,
            max_row=5
        )

        category_ref = Reference(
            ws,
            min_col=1,
            min_row=2,
            max_row=5
        )

        chart.add_data(
            data_ref,
            titles_from_data=True
        )

        chart.set_categories(
            category_ref
        )

        ws.add_chart(
            chart,
            "F2"
        )

        # Conditional Formatting
        red_fill = PatternFill(
            start_color="FF9999",
            end_color="FF9999",
            fill_type="solid"
        )

        ws.conditional_formatting.add(
            "D2:D5",
            CellIsRule(
                operator="lessThan",
                formula=["200000"],
                fill=red_fill
            )
        )

        wb.save(file_name)

        print(
            "Workbook created successfully."
        )

    @staticmethod
    def read_workbook(file_name):

        wb = load_workbook(file_name)

        print(
            "\nWorkbook Sheets:",
            wb.sheetnames
        )

        for sheet in wb.sheetnames:

            ws = wb[sheet]

            print(
                f"\nSheet: {sheet}"
            )

            for row in ws.iter_rows(
                values_only=True
            ):

                print(row)


# ==========================================
# DEMO
# ==========================================

if __name__ == "__main__":

    sample_data = [

        {
            "ID": "1",
            "Name": "Suresh"
        },

        {
            "ID": "2",
            "Name": "Ravi"
        },

        {
            "ID": "2",
            "Name": "Ravi"
        }

    ]

    # Write CSV
    CSVProcessor.write_csv(
        "employees.csv",
        sample_data
    )

    # Read CSV
    data = CSVProcessor.read_csv(
        "employees.csv"
    )

    # Remove Duplicates
    unique_data = (
        CSVProcessor.remove_duplicates(
            data,
            "ID"
        )
    )

    print("\nUnique Records:")
    print(unique_data)

    # Validate CSV
    errors = (
        CSVProcessor.validate_data(
            data,
            ["ID", "Name"]
        )
    )

    print("\nValidation Errors:")
    print(errors)

    # Create Excel Workbook
    ExcelProcessor.create_workbook(
        "sales_report.xlsx"
    )

    # Read Workbook
    ExcelProcessor.read_workbook(
        "sales_report.xlsx"
    )